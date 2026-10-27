from __future__ import annotations

from dataclasses import replace
import inspect
import hashlib
import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.cell.rich_text import CellRichText
from openpyxl.styles import PatternFill

from src import main as main_module
from src.inspection_adjustment import (
    SYSTEM_YELLOW,
    TARGET_HEADERS,
    InspectionAdjustmentError,
)
from src.inspection_adjustment_excel import (
    InspectionAdjustmentService,
    create_empty_target_workbook,
)
from src.ui import IqcApp
from src.ui_preferences import UiPreferenceStore
from tests.fixtures.inspection_adjustment_fixture import (
    empty_target_row,
    source_row,
    write_source_workbook,
    write_target_workbook,
)


class InspectionAdjustmentServiceTests(unittest.TestCase):
    def setUp(self) -> None:
        self.temp = tempfile.TemporaryDirectory()
        self.root = Path(self.temp.name)
        self.source = self.root / "a.xlsx"
        self.target = self.root / "b.xlsx"
        self.service = InspectionAdjustmentService()

    def tearDown(self) -> None:
        self.temp.cleanup()

    def test_preview_maps_values_sorts_only_new_batch_and_keeps_unmapped_blank(self) -> None:
        write_source_workbook(
            self.source,
            [
                source_row("DEMO-QM-0018"),
                source_row("DEMO-QM-0016", purchase="PO-002", quantity=3),
            ],
        )
        historical = empty_target_row()
        historical[0] = "历史送检"
        historical[19] = "DEMO-QM-0015"
        write_target_workbook(self.target, rows=[tuple(historical)])

        preview = self.service.build_preview(self.source, self.target)

        self.assertEqual(preview.new_count, 2)
        self.assertEqual(
            [record.identifier for record in preview.new_records],
            ["DEMO-QM-0016", "DEMO-QM-0018"],
        )
        values = preview.new_records[0].target_values
        self.assertEqual(values[0], "IN-001")
        self.assertEqual(values[1], "PO-002")
        self.assertEqual(values[2], None)
        self.assertEqual(values[3], "M001")
        self.assertEqual(values[4], "物料一")
        self.assertEqual(values[5], "供应商一")
        self.assertEqual(values[9], 3)
        self.assertEqual(values[19], "DEMO-QM-0016")

    def test_sync_updates_original_and_marks_only_new_rows(self) -> None:
        write_source_workbook(
            self.source,
            [
                source_row("DEMO-QM-0018"),
                source_row("DEMO-QM-0016"),
            ],
        )
        historical = empty_target_row()
        historical[0] = "历史送检"
        historical[19] = "DEMO-QM-0015"
        write_target_workbook(self.target, rows=[tuple(historical)])

        preview = self.service.build_preview(self.source, self.target)
        result = self.service.synchronize(preview)

        self.assertTrue(result.written)
        self.assertEqual(result.new_rows, (
            ("DEMO-QM-0016", 3),
            ("DEMO-QM-0018", 4),
        ))
        self.assertIsNone(result.backup_path)
        self.assertEqual(list(self.target.parent.glob("*.backup.*.xlsx")), [])
        workbook = load_workbook(self.target, data_only=True)
        try:
            sheet = workbook["IQC统计"]
            self.assertEqual(sheet["T2"].value, "DEMO-QM-0015")
            self.assertEqual(sheet["T3"].value, "DEMO-QM-0016")
            self.assertEqual(sheet["T4"].value, "DEMO-QM-0018")
            self.assertEqual(sheet["C3"].value, None)
            self.assertEqual(sheet["Q3"].value, None)
            self.assertEqual(sheet["T3"].fill.fgColor.rgb[-6:].upper(), SYSTEM_YELLOW)
            self.assertEqual(sheet["T4"].fill.fgColor.rgb[-6:].upper(), SYSTEM_YELLOW)
        finally:
            workbook.close()

    def test_sync_highlights_only_item_code_suffix_in_target_a(self) -> None:
        write_source_workbook(
            self.source,
            [
                source_row(
                    "DEMO-QM-0004",
                    incoming_number="RK9999000000000004",
                ),
                source_row(
                    "DEMO-QM-0005",
                    incoming_number="RK9999X",
                ),
            ],
        )
        write_target_workbook(self.target)

        preview = self.service.build_preview(self.source, self.target)
        result = self.service.synchronize(preview)

        self.assertTrue(result.written)
        # 默认 data_only 回读保持完整普通字符串，去重仍按原值工作。
        workbook = load_workbook(self.target, data_only=True)
        try:
            sheet = workbook["IQC统计"]
            self.assertEqual(sheet["A2"].value, "RK9999000000000004")
            self.assertEqual(sheet["A3"].value, "RK9999X")
            self.assertEqual(sheet["T2"].value, "DEMO-QM-0004")
        finally:
            workbook.close()

        # rich_text=True 仅用于核对显示分段；前缀保持普通文本，后十位蓝色粗体。
        workbook = load_workbook(self.target, data_only=False, rich_text=True)
        try:
            sheet = workbook["IQC统计"]
            item_code = sheet["A2"].value
            self.assertIsInstance(item_code, CellRichText)
            prefix = list(item_code)[0]
            self.assertEqual(prefix.text, "RK999900")
            self.assertEqual(prefix.font.rFont, "微软雅黑")
            self.assertEqual(prefix.font.sz, 11.0)
            suffix = list(item_code)[1]
            self.assertEqual(suffix.text, "0000000004")
            self.assertEqual(suffix.font.rFont, "微软雅黑")
            self.assertEqual(suffix.font.sz, 11.0)
            self.assertTrue(suffix.font.b)
            self.assertEqual(suffix.font.color.type, "rgb")
            self.assertEqual(suffix.font.color.rgb[-6:].upper(), "0070C0")
            self.assertEqual(sheet["A2"].font.name, "微软雅黑")
            self.assertEqual(sheet["A2"].font.sz, 11.0)
            self.assertNotIsInstance(sheet["A3"].value, CellRichText)
            self.assertEqual(sheet["A3"].font.name, "微软雅黑")
            self.assertEqual(sheet["A3"].font.sz, 11.0)
            self.assertNotIsInstance(sheet["T2"].value, CellRichText)
        finally:
            workbook.close()

        # 第二次读取会将同一物项编码判为重复，证明富文本不改变去重语义。
        reread = self.service.build_preview(self.source, self.target)
        self.assertEqual(reread.new_count, 0)
        self.assertEqual(reread.duplicate_count, 2)

        # 再追加一条时保留历史富文本，且新行继续按同一规则写入。
        write_source_workbook(
            self.source,
            [
                source_row(
                    "DEMO-QM-0004",
                    incoming_number="RK9999000000000004",
                ),
                source_row(
                    "DEMO-QM-0005",
                    incoming_number="RK9999X",
                ),
                source_row(
                    "DEMO-QM-0006",
                    incoming_number="RK9999000000000006",
                ),
            ],
        )
        second_preview = self.service.build_preview(self.source, self.target)
        self.assertEqual(second_preview.new_count, 1)
        self.assertEqual(second_preview.duplicate_count, 2)
        self.service.synchronize(second_preview)
        workbook = load_workbook(self.target, data_only=False, rich_text=True)
        try:
            sheet = workbook["IQC统计"]
            self.assertIsInstance(sheet["A2"].value, CellRichText)
            self.assertIsInstance(sheet["A4"].value, CellRichText)
            self.assertEqual(sheet["A4"].value[-1].text, "0000000006")
        finally:
            workbook.close()

    def test_all_duplicate_records_do_not_write_or_create_backup(self) -> None:
        source_values = source_row("DEMO-QM-0016")
        write_source_workbook(self.source, [source_values])
        existing = empty_target_row()
        existing[19] = "DEMO-QM-0016"
        write_target_workbook(self.target, rows=[tuple(existing)])
        before = self.target.stat().st_mtime_ns

        preview = self.service.build_preview(self.source, self.target)
        result = self.service.synchronize(preview)

        self.assertEqual(preview.new_count, 0)
        self.assertEqual(preview.duplicate_count, 1)
        self.assertFalse(result.written)
        self.assertIsNone(result.backup_path)
        self.assertEqual(self.target.stat().st_mtime_ns, before)
        self.assertEqual(list(self.target.parent.glob("*.backup.*.xlsx")), [])

    def test_conflicting_duplicate_identifier_blocks_sync_without_backup_or_fill_changes(self) -> None:
        write_source_workbook(
            self.source,
            [
                source_row("DEMO-QM-0016", purchase="PO-001"),
                source_row("DEMO-QM-0016", purchase="PO-002"),
            ],
        )
        existing = empty_target_row()
        existing[0] = "历史送检"
        existing[19] = "DEMO-QM-0015"
        write_target_workbook(self.target, rows=[tuple(existing)])
        workbook = load_workbook(self.target)
        try:
            sheet = workbook["IQC统计"]
            for column in range(1, 21):
                sheet.cell(row=2, column=column).fill = PatternFill(fill_type="solid", fgColor=SYSTEM_YELLOW)
            workbook.save(self.target)
        finally:
            workbook.close()
        before_bytes = hashlib.sha256(self.target.read_bytes()).hexdigest()

        preview = self.service.build_preview(self.source, self.target)

        self.assertEqual(preview.new_count, 0)
        self.assertEqual(preview.conflict_identifiers, ("DEMO-QM-0016",))
        self.assertEqual(preview.anomaly_count, 1)
        with self.assertRaises(InspectionAdjustmentError) as context:
            self.service.synchronize(preview)
        self.assertIn("人工核查", str(context.exception))
        self.assertIn("字段冲突", str(context.exception))
        self.assertEqual(hashlib.sha256(self.target.read_bytes()).hexdigest(), before_bytes)
        self.assertEqual(list(self.target.parent.glob("*.backup.*.xlsx")), [])
        workbook = load_workbook(self.target, data_only=True)
        try:
            sheet = workbook["IQC统计"]
            self.assertEqual(sheet["T2"].value, "DEMO-QM-0015")
            self.assertEqual(sheet["A3"].value, None)
            self.assertEqual(sheet["A2"].fill.fgColor.rgb[-6:].upper(), SYSTEM_YELLOW)
        finally:
            workbook.close()

    def test_empty_and_invalid_identifiers_are_anomalies(self) -> None:
        write_source_workbook(
            self.source,
            [
                source_row(""),
                source_row("#VALUE!"),
                source_row("DEMO-QM-0016"),
            ],
        )
        write_target_workbook(self.target)

        preview = self.service.build_preview(self.source, self.target)

        self.assertEqual(preview.scanned_count, 3)
        self.assertEqual(preview.valid_count, 1)
        self.assertEqual(preview.anomaly_count, 2)
        self.assertEqual(preview.new_count, 1)

    def test_changed_target_after_preview_requires_reread(self) -> None:
        write_source_workbook(self.source, [source_row("DEMO-QM-0016")])
        write_target_workbook(self.target)
        preview = self.service.build_preview(self.source, self.target)

        with self.target.open("ab") as stream:
            stream.write(b"changed-after-preview")

        with self.assertRaises(InspectionAdjustmentError) as context:
            self.service.synchronize(preview)
        self.assertTrue(context.exception.requires_reread)

    def test_same_actual_file_is_rejected_at_preview_and_sync_entry(self) -> None:
        write_target_workbook(self.source)

        with self.assertRaises(InspectionAdjustmentError) as context:
            self.service.build_preview(self.source, self.source)
        self.assertFalse(context.exception.retryable)
        self.assertIn("不允许 A/B 指向同一文件", str(context.exception))
        self.assertIn("未修改 A", str(context.exception))

        write_source_workbook(self.source, [source_row("DEMO-QM-0016")])
        write_target_workbook(self.target)
        preview = self.service.build_preview(self.source, self.target)
        same_file_preview = replace(preview, source_path=self.target)
        with self.assertRaises(InspectionAdjustmentError) as context:
            self.service.synchronize(same_file_preview)
        self.assertFalse(context.exception.retryable)
        self.assertIn("不允许 A/B 指向同一文件", str(context.exception))

    def test_display_semantics_preserve_formatted_numeric_identifiers_and_mapped_codes(self) -> None:
        write_source_workbook(
            self.source,
            [
                source_row(15, item_code=42),
                source_row(16, item_code=43),
            ],
        )
        existing = empty_target_row()
        existing[19] = 15
        write_target_workbook(self.target, rows=[tuple(existing)])

        source_workbook = load_workbook(self.source)
        try:
            source_sheet = source_workbook["质检信息导出-1"]
            source_sheet["D2"].number_format = "000000"
            source_sheet["D3"].number_format = "000000"
            source_sheet["G2"].number_format = "000000"
            source_sheet["G3"].number_format = "000000"
            source_workbook.save(self.source)
        finally:
            source_workbook.close()
        target_workbook = load_workbook(self.target)
        try:
            target_workbook["IQC统计"]["T2"].number_format = "000000"
            target_workbook.save(self.target)
        finally:
            target_workbook.close()

        preview = self.service.build_preview(self.source, self.target)

        self.assertEqual(preview.duplicate_count, 1)
        self.assertEqual(preview.new_count, 1)
        self.assertEqual(preview.new_records[0].identifier, "000016")
        self.assertEqual(preview.new_records[0].target_values[3], "000043")

    def test_final_fingerprint_change_blocks_and_preserves_temp_without_backup(self) -> None:
        write_source_workbook(self.source, [source_row("DEMO-QM-0016")])
        write_target_workbook(self.target)
        preview = self.service.build_preview(self.source, self.target)

        original_verify = self.service._verify_written_workbook
        verification_calls = 0

        def verify(path, preview_arg, new_rows):
            nonlocal verification_calls
            original_verify(path, preview_arg, new_rows)
            verification_calls += 1
            if verification_calls == 1:
                workbook = load_workbook(self.target)
                try:
                    workbook["IQC统计"]["A2"] = "外部修改"
                    workbook.save(self.target)
                finally:
                    workbook.close()

        self.service._verify_written_workbook = verify
        with self.assertRaises(InspectionAdjustmentError) as context:
            self.service.synchronize(preview)

        self.assertTrue(context.exception.requires_reread)
        self.assertIn("临时文件", str(context.exception))
        self.assertEqual(list(self.target.parent.glob("*.backup.*.xlsx")), [])
        self.assertTrue(list(self.target.parent.glob(".*.tmp.xlsx")))

    def test_post_replace_readback_failure_does_not_create_backup(self) -> None:
        write_source_workbook(self.source, [source_row("DEMO-QM-0016")])
        write_target_workbook(self.target)
        preview = self.service.build_preview(self.source, self.target)
        original_verify = self.service._verify_written_workbook
        verification_calls = 0

        def verify(path, preview_arg, new_rows):
            nonlocal verification_calls
            verification_calls += 1
            if verification_calls == 2:
                raise InspectionAdjustmentError("模拟替换后回读失败")
            original_verify(path, preview_arg, new_rows)

        self.service._verify_written_workbook = verify
        with self.assertRaises(InspectionAdjustmentError):
            self.service.synchronize(preview)

        self.assertEqual(list(self.target.parent.glob("*.backup.*.xlsx")), [])
        workbook = load_workbook(self.target, data_only=True)
        try:
            self.assertEqual(workbook["IQC统计"]["T2"].value, "DEMO-QM-0016")
        finally:
            workbook.close()
        self.assertEqual(list(self.target.parent.glob(".*.restore.*.tmp.xlsx")), [])

    def test_missing_source_sheet_and_multiple_target_sheets_block(self) -> None:
        workbook = Workbook()
        workbook.active.title = "wrong"
        workbook.save(self.source)
        workbook.close()
        write_target_workbook(self.target)
        with self.assertRaises(InspectionAdjustmentError):
            self.service.build_preview(self.source, self.target)

        write_source_workbook(self.source, [source_row("DEMO-QM-0016")])
        workbook = load_workbook(self.target)
        workbook.create_sheet("extra")
        workbook.save(self.target)
        workbook.close()
        with self.assertRaises(InspectionAdjustmentError):
            self.service.build_preview(self.source, self.target)

    def test_create_empty_target_uses_one_iqc_sheet_and_contract_headers(self) -> None:
        create_empty_target_workbook(self.target)
        workbook = load_workbook(self.target, data_only=True)
        try:
            self.assertEqual(workbook.sheetnames, ["IQC统计"])
            self.assertEqual(
                tuple(workbook["IQC统计"].cell(row=1, column=index).value for index in range(1, 21)),
                tuple(value or None for value in TARGET_HEADERS),
            )
        finally:
            workbook.close()

    def test_create_empty_target_refuses_to_overwrite_existing_file(self) -> None:
        write_target_workbook(self.target)
        before_bytes = hashlib.sha256(self.target.read_bytes()).hexdigest()

        with self.assertRaises(InspectionAdjustmentError) as context:
            create_empty_target_workbook(self.target)

        self.assertFalse(context.exception.retryable)
        self.assertIn("拒绝覆盖", str(context.exception))
        self.assertEqual(
            hashlib.sha256(self.target.read_bytes()).hexdigest(),
            before_bytes,
        )


class InspectionAdjustmentPreferenceTests(unittest.TestCase):
    def test_a_and_b_paths_are_independent_and_unknown_fields_survive(self) -> None:
        with tempfile.TemporaryDirectory() as temp:
            path = Path(temp) / "ui_preferences.json"
            store = UiPreferenceStore(path)
            path.write_text('{"last_active_module":"IQC","custom":"keep"}', encoding="utf-8")

            self.assertEqual(store.save_inspection_adjustment_a_path("A:/a.xlsx"), (True, ""))
            self.assertEqual(store.save_inspection_adjustment_b_path("B:/b.xlsx"), (True, ""))
            loaded = store.load()

            self.assertEqual(loaded.last_inspection_adjustment_a_path, "A:/a.xlsx")
            self.assertEqual(loaded.last_inspection_adjustment_b_path, "B:/b.xlsx")
            self.assertEqual(store.save("INSPECTION_ADJUSTMENT"), (True, ""))
            self.assertEqual(store.load().module, "INSPECTION_ADJUSTMENT")
            self.assertIn('"custom": "keep"', path.read_text(encoding="utf-8"))


class InspectionAdjustmentAssemblyTests(unittest.TestCase):
    def test_main_injects_service_without_external_business_dependencies(self) -> None:
        source = inspect.getsource(main_module.main)
        self.assertIn("InspectionAdjustmentService()", source)
        self.assertIn("inspection_adjustment_service=inspection_adjustment_service", source)
        self.assertNotIn("run_generated_sql.bat", source.split("inspection_adjustment_service =")[1].split("app = IqcApp", 1)[0])

    def test_set_busy_covers_inspection_adjustment_tab_and_page(self) -> None:
        source = inspect.getsource(IqcApp._set_busy)
        self.assertIn("inspection_adjustment_tab_button", source)
        self.assertIn("inspection_adjustment_page", source)
        self.assertIn("set_busy(busy)", source)

    def test_build_passes_loaded_preference_to_third_page_without_local_reference(self) -> None:
        init_source = inspect.getsource(IqcApp.__init__)
        build_source = inspect.getsource(IqcApp._build)

        self.assertIn("self._preference = preference", init_source)
        self.assertIn("InspectionAdjustmentPage(", build_source)
        self.assertIn("preference=self._preference", build_source)
        self.assertNotIn("preference=preference", build_source)
        self.assertNotIn("preference", IqcApp._build.__code__.co_varnames)


if __name__ == "__main__":
    unittest.main()
