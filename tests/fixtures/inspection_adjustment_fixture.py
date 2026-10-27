"""送检信息调整的脱敏临时工作簿构造器。

夹具只写入调用方提供的临时目录，不包含真实 UNC 路径、真实供应商、
数据库、Windchill 或共享工作簿内容。
"""

from __future__ import annotations

from pathlib import Path
from typing import Iterable, Mapping

from openpyxl import Workbook

from src.inspection_adjustment import SOURCE_SHEET_NAME, TARGET_HEADERS


def write_source_workbook(path: Path, rows: Iterable[Mapping[str, object]]) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = SOURCE_SHEET_NAME
    for row_number, row in enumerate(rows, start=2):
        for column, value in row.items():
            sheet[f"{column}{row_number}"] = value
    workbook.save(path)
    workbook.close()


def write_target_workbook(
    path: Path,
    *,
    rows: Iterable[Iterable[object]] = (),
    sheet_name: str = "IQC统计",
) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = sheet_name
    for column, value in enumerate(TARGET_HEADERS, start=1):
        sheet.cell(row=1, column=column, value=value or None)
    for row_number, values in enumerate(rows, start=2):
        for column, value in enumerate(values, start=1):
            sheet.cell(row=row_number, column=column, value=value)
    workbook.save(path)
    workbook.close()


def source_row(
    identifier: object,
    *,
    purchase: object = "PO-001",
    quantity: object = 2,
    item_code: object = "M001",
    item_name: object = "物料一",
    supplier: object = "供应商一",
    incoming_number: object = "IN-001",
) -> dict[str, object]:
    return {
        "A": purchase,
        "C": incoming_number,
        "D": identifier,
        "G": item_code,
        "H": item_name,
        "L": quantity,
        "Q": supplier,
    }


def empty_target_row() -> list[object]:
    return [None] * len(TARGET_HEADERS)
